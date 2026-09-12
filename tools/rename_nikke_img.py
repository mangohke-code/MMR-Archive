# -*- coding: utf-8 -*-
"""니케 이미지 파일을 게임 코드 이름으로 바꾼다.

왜 코드로 바꾸나
  한글 이름과 코드 이름(mi_cXXX_YY_s.png)이 반반 섞여 있어서, 같은 캐릭터 파일이
  두 번 올라가도 이름이 안 겹쳐 아무것도 안 걸렸다. 코드로 통일하면 파일 이름
  자체가 중복 방지 장치가 된다.

왜 옮기지 않고 복사하나
  깃허브 페이지와 DB 는 따로 배포된다. 파일을 옮겨 버리면 DB 가 아직 옛 이름을
  가리키는 동안 이미지가 통째로 깨진다. 새 이름을 먼저 "추가" 해서 둘 다 있는
  상태를 만들고, SQL 을 돌린 뒤에 옛 파일을 지우면 끊기는 순간이 없다.

png 는 webp 로 같이 바꾼다 - 같은 그림이 77% 작아진다(재서 확인).
"""
import io
import os
import re
import sys

from openpyxl import load_workbook
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
WEB = os.path.dirname(HERE)
XLSX = os.path.join(os.path.dirname(WEB), '니케 이미지 코드.xlsx')
SRC_DIR = os.path.join(WEB, 'img', 'nikke')
QUALITY = 90


def plan():
    ws = load_workbook(XLSX, data_only=True)['코드 입력']
    out = []
    for r in range(3, ws.max_row + 1):
        name, no, col, cur, code = [ws.cell(row=r, column=c).value for c in range(1, 6)]
        if not name or not code:
            continue
        code = str(code).strip()
        if not re.fullmatch(r'c\d{3,4}_\d{2}', code):
            raise SystemExit('코드 형식이 이상하다: %r (%s)' % (code, name))
        out.append((no, col, name, str(cur).strip(), code + '.webp'))
    return out


def main():
    rows = plan()
    made = same = 0
    for no, col, name, cur, new in rows:
        src = os.path.join(SRC_DIR, cur)
        dst = os.path.join(SRC_DIR, new)
        if not os.path.exists(src):
            raise SystemExit('원본이 없다: %s' % src)
        if os.path.exists(dst):
            same += 1
            continue
        if cur.lower().endswith('.webp'):
            # 이미 webp 면 다시 인코딩하지 않는다. 손실이 한 번 더 얹힌다.
            with open(src, 'rb') as f:
                data = f.read()
            with open(dst, 'wb') as f:
                f.write(data)
        else:
            im = Image.open(src).convert('RGBA')
            im.save(dst, 'WEBP', quality=QUALITY, method=6)
        made += 1
    print('새로 만든 파일 %d개 / 이미 있던 것 %d개' % (made, same))

    # DB 갱신 SQL
    sql = io.open(os.path.join(os.path.dirname(WEB), 'set_nikke_img_code.sql'), 'w', encoding='utf-8')
    sql.write('-- 니케 이미지 경로를 코드 이름으로 바꾼다.\n')
    sql.write('-- tools/rename_nikke_img.py 가 새 이름 파일을 이미 만들어 뒀다(옛 파일도 그대로 있다).\n')
    sql.write('-- 순서: 새 파일 푸시 -> 배포 확인 -> 이 SQL -> 화면 확인 -> 옛 파일 삭제\n\n')
    for no, col, name, cur, new in rows:
        sql.write('update "IMG_니케" set "%s" = \'img/nikke/%s\' where "번호" = %s;  -- %s (%s)\n'
                  % (col, new, no, name, cur))
    sql.write('\n-- 확인용: 아직 옛 이름을 가리키는 행\n')
    sql.write('select "번호", "이름", "이미지", "코스튬1_이미지", "코스튬2_이미지"\n')
    sql.write('  from "IMG_니케"\n')
    sql.write(' where "이미지" not like \'img/nikke/c%\' and "이미지" not like \'http%\'\n')
    sql.write('    or "코스튬1_이미지" not like \'img/nikke/c%\'\n')
    sql.write('    or "코스튬2_이미지" not like \'img/nikke/c%\';\n')
    sql.write('\n-- 위 SQL 을 돌리고 화면을 확인한 뒤 지워도 되는 옛 파일 %d개:\n' % len(rows))
    for no, col, name, cur, new in sorted(rows, key=lambda x: x[3]):
        sql.write('--   img/nikke/%s\n' % cur)
    sql.close()
    print('SQL: set_nikke_img_code.sql (%d줄)' % len(rows))


if __name__ == '__main__':
    main()
